import datetime as dt
import json
import math
import pathlib
import re

import openpyxl


SOURCE = pathlib.Path(r"\\Danieta_Nas\Danieta\Casa\Facturas.xlsx")
ROOT = pathlib.Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "seed.json"
PUBLIC_DATA_PATH = ROOT / "public" / "facturas" / "data.json"


def clean_text(value):
    if value is None:
        return None
    return str(value).replace("�", "ó").strip()


def number(value):
    if isinstance(value, (int, float)) and math.isfinite(value):
        return round(float(value), 6)
    return None


def iso_date(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    if not value:
        return None
    text = str(value).strip()
    match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", text)
    if match:
        day, month, year = map(int, match.groups())
        return dt.date(year, month, day).isoformat()
    try:
        return dt.datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return None


book = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)

summary = []
sheet = book["Resumen Hist"]
property_name = "Total"
for row in range(1, sheet.max_row + 1):
    label = clean_text(sheet.cell(row, 1).value)
    if not label:
        continue
    if sheet.cell(row, 2).value == "Inicio":
        property_name = label.rsplit(" - ", 1)[0]
        continue
    summary.append(
        {
            "property": property_name,
            "category": label,
            "start_date": iso_date(sheet.cell(row, 2).value),
            "end_date": iso_date(sheet.cell(row, 3).value),
            "days": int(sheet.cell(row, 4).value or 0),
            "daily_average": number(sheet.cell(row, 5).value),
            "total": number(sheet.cell(row, 6).value),
        }
    )

sheet = book["Facturas"]
invoice_columns = {
    "Alquiler": (3, 4, 7),
    "Luz": (8, 10, 13),
    "Calefacción": (14, 16, 19),
    "Agua": (20, 22, 25),
    "Internet": (26, 27, 30),
    "Limpieza": (31, 32, 35),
}
invoices = []
property_name = None
for row in range(3, sheet.max_row + 1):
    house = clean_text(sheet.cell(row, 2).value)
    if house and house != "Total":
        property_name = house
    if house == "Total" or not property_name:
        continue
    for category, (start_col, end_col, amount_col) in invoice_columns.items():
        start_date = iso_date(sheet.cell(row, start_col).value)
        end_date = iso_date(sheet.cell(row, end_col).value)
        amount = number(sheet.cell(row, amount_col).value)
        if not start_date or not end_date or amount is None:
            continue
        calculated_days = (dt.date.fromisoformat(end_date) - dt.date.fromisoformat(start_date)).days + 1
        if calculated_days < 1:
            continue
        invoices.append(
            {
                "id": f"excel-facturas-{row}-{category.lower().replace('ó', 'o')}",
                "source": "excel",
                "source_sheet": "Facturas",
                "source_row": row,
                "property": property_name,
                "category": category,
                "start_date": start_date,
                "end_date": end_date,
                "days": calculated_days,
                "daily_average": round(amount / calculated_days, 6),
                "total": amount,
            }
        )


monthly = []
monthly_map = {}
for invoice in invoices:
    key = (invoice["property"], invoice["start_date"][:7] + "-01")
    if key not in monthly_map:
        monthly_map[key] = {
            "property": invoice["property"],
            "month": key[1],
            "daily_average": 0,
            "total": 0,
            "categories": {category: 0 for category in invoice_columns},
            "category_daily": {category: 0 for category in invoice_columns},
            "category_details": {category: [] for category in invoice_columns},
        }
    item = monthly_map[key]
    category = invoice["category"]
    item["categories"][category] += invoice["total"]
    item["category_daily"][category] += invoice["daily_average"]
    item["category_details"][category].append(
        {
            "start_date": invoice["start_date"],
            "end_date": invoice["end_date"],
            "days": invoice["days"],
            "daily_average": invoice["daily_average"],
            "total": invoice["total"],
        }
    )

for item in monthly_map.values():
    item["categories"] = {category: round(value, 6) for category, value in item["categories"].items()}
    item["category_daily"] = {category: round(value, 6) for category, value in item["category_daily"].items()}
    item["total"] = round(sum(item["categories"].values()), 6)
    item["daily_average"] = round(sum(item["category_daily"].values()), 6)
    monthly.append(item)
monthly.sort(key=lambda item: (item["month"], item["property"]))

electricity = []
sheet = book["Consumo Luz"]
for row in range(2, sheet.max_row + 1):
    start_date = iso_date(sheet.cell(row, 4).value)
    end_date = iso_date(sheet.cell(row, 5).value)
    amount = number(sheet.cell(row, 8).value)
    if not start_date or not end_date or amount is None:
        continue
    days = (dt.date.fromisoformat(end_date) - dt.date.fromisoformat(start_date)).days + 1
    p1 = number(sheet.cell(row, 9).value) or 0
    p2 = number(sheet.cell(row, 12).value) or 0
    p3 = number(sheet.cell(row, 15).value) or 0
    consumption = round(p1 + p2 + p3, 6)
    electricity.append(
        {
            "property": clean_text(sheet.cell(row, 1).value),
            "company": clean_text(sheet.cell(row, 2).value),
            "month": start_date[:7] + "-01",
            "start_date": start_date,
            "end_date": end_date,
            "days": days,
            "amount": amount,
            "p1": p1,
            "p2": p2,
            "p3": p3,
            "consumption": consumption,
            "daily_consumption": round(consumption / days, 6),
            "daily_amount": round(amount / days, 6),
            "cost_per_kwh": round(amount / consumption, 6) if consumption else 0,
            "p1_share": round(p1 / consumption, 6) if consumption else 0,
            "p2_share": round(p2 / consumption, 6) if consumption else 0,
            "p3_share": round(p3 / consumption, 6) if consumption else 0,
        }
    )

floor = []
sheet = book["Consumo Suelo"]
for row in range(2, sheet.max_row + 1):
    start_date = iso_date(sheet.cell(row, 4).value)
    end_date = iso_date(sheet.cell(row, 5).value)
    if not start_date or not end_date:
        continue
    days = (dt.date.fromisoformat(end_date) - dt.date.fromisoformat(start_date)).days + 1
    cold_kwh = number(sheet.cell(row, 9).value) or 0
    cold_price = number(sheet.cell(row, 10).value) or 0
    heat_kwh = number(sheet.cell(row, 14).value) or 0
    heat_price = number(sheet.cell(row, 15).value) or 0
    water_m3 = number(sheet.cell(row, 19).value) or 0
    water_price = number(sheet.cell(row, 20).value) or 0
    fixed = number(sheet.cell(row, 24).value) or 0
    cold_total = round(cold_kwh * cold_price, 6)
    heat_total = round(heat_kwh * heat_price, 6)
    water_total = round(water_m3 * water_price, 6)
    amount = round((cold_total + heat_total + water_total + fixed) * 1.21, 6)
    floor.append(
        {
            "property": clean_text(sheet.cell(row, 1).value),
            "company": clean_text(sheet.cell(row, 2).value),
            "month": start_date[:7] + "-01",
            "start_date": start_date,
            "end_date": end_date,
            "days": days,
            "amount": amount,
            "cold_kwh": cold_kwh,
            "heat_kwh": heat_kwh,
            "water_m3": water_m3,
            "fixed": fixed,
            "cold_price": cold_price,
            "cold_total": cold_total,
            "heat_price": heat_price,
            "heat_total": heat_total,
            "water_price": water_price,
            "water_total": water_total,
            "daily_amount": round(amount / days, 6),
            "cold_share": round(cold_total / amount, 6) if amount else 0,
            "heat_share": round(heat_total / amount, 6) if amount else 0,
            "water_share": round(water_total / amount, 6) if amount else 0,
            "fixed_share": round(fixed / amount, 6) if amount else 0,
            "vat_rate": 21,
        }
    )

vehicle = []
sheet = book["Gastos Coche"]
for row in range(2, sheet.max_row + 1):
    date = iso_date(sheet.cell(row, 1).value)
    price_or_amount = number(sheet.cell(row, 4).value)
    extended_amount = number(sheet.cell(row, 5).value)
    amount = extended_amount if extended_amount is not None else price_or_amount
    if date and amount is not None:
        description = clean_text(sheet.cell(row, 2).value)
        is_fuel_table = extended_amount is not None
        vehicle.append(
            {
                "date": date,
                "description": description,
                "category": "Gasolina",
                "vehicle": clean_text(sheet.cell(row, 3).value),
                "amount": amount,
                "unit_price": price_or_amount if is_fuel_table else None,
                "quantity": number(sheet.cell(row, 6).value) if is_fuel_table else None,
            }
        )
    service_date = iso_date(sheet.cell(row, 14).value)
    service_amount = number(sheet.cell(row, 21).value)
    if service_date and service_amount is not None:
        vehicle.append(
            {
                "date": service_date,
                "description": clean_text(sheet.cell(row, 16).value),
                "category": clean_text(sheet.cell(row, 20).value) or "Otros",
                "vehicle": None,
                "amount": service_amount,
            }
        )

payload = {
    "meta": {
        "source_updated": SOURCE.stat().st_mtime,
        "generated": dt.datetime.now(dt.timezone.utc).isoformat(),
        "excluded_sheets": ["Previsión Luz", "X Gasolina Moto", "Comparación"],
    },
    "summary": summary,
    "invoices": sorted(invoices, key=lambda item: (item["start_date"], item["property"], item["category"])),
    "monthly": monthly,
    "electricity": electricity,
    "floor": floor,
    "vehicle": sorted(vehicle, key=lambda item: item["date"]),
}

DATA_PATH.parent.mkdir(parents=True, exist_ok=True)
DATA_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
PUBLIC_DATA_PATH.parent.mkdir(parents=True, exist_ok=True)
PUBLIC_DATA_PATH.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

print(
    json.dumps(
        {key: len(value) for key, value in payload.items() if isinstance(value, list)},
        ensure_ascii=False,
    )
)
